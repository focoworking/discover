"""Exporta el estudio a producto/funnel26/datos.js del sitio FOCO.
Uso: python3 research/scripts/export_funnel26.py /ruta/a/foco/producto/funnel26/datos.js
Lee valores calculados del Excel (puntaje y prioridad) y los módulos de data/."""
import json, os, sys
from datetime import date
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "data"))
from macro import AREAS, COUNTIES, DIGITAL
from ads_legal import ADS, PRECIOS, LEGAL
from registros import REGISTROS
from coordenadas import COORDS

out = sys.argv[1]
ws = load_workbook(os.path.join(ROOT, "Miami_Oportunidad_Digital.xlsx"), data_only=True)["Prospectos"]
hdr = [c.value for c in ws[4]]
key = {"ID":"id","Empresa":"nombre","Municipio":"municipio","Barrio/zona":"barrio","Condado":"condado","Sector (detalle)":"sector",
       "NAICS":"naics","Industria (NAICS)":"industria","Forma legal":"legal","Estado web":"web","Redes sociales":"redes",
       "Tiene redes":"tieneRedes","Evidencia":"evidencia","Fuente":"fuente","Empleados (est.)":"empleados","Año fundación":"fundada",
       "Puntaje":"puntaje","Prioridad":"prioridad","Estado Sunbiz":"sunbiz","Alerta":"alerta","Nicho":"nicho"}
CKEY = {"Dirección":"direccion","Teléfono":"telefono","Web":"web","Facebook":"facebook","Instagram":"instagram",
        "LinkedIn":"linkedin","Yelp":"yelp","Fuente contacto":"fuente","Nota contacto":"nota"}
pros = []
for row in ws.iter_rows(min_row=5, values_only=True):
    if not row[0]: continue
    d = {key[h]: v for h, v in zip(hdr, row) if h in key}
    d = {k: v for k, v in d.items() if v not in (None, "")}
    c = {CKEY[h]: v for h, v in zip(hdr, row) if h in CKEY and v not in (None, "")}
    if c: d["contacto"] = c
    pros.append(d)

datos = {
  "corte": date.today().isoformat(),
  "prospectos": pros,
  "areas": [dict(zip(["area","condado","poblacion","anioPob","firmas","caracter","naics","verificacion","fuente"], a)) for a in AREAS],
  "condados": [dict(zip(["condado","poblacion","establecimientos","anio","est20a99","est100a499","fuente"], c)) for c in COUNTIES],
  "digital": [dict(zip(["indicador","valor","anio","origen","fuente"], d)) for d in DIGITAL],
  "ads": [dict(zip(["industria","naics","cpc","cpl","nota","fuente"], a)) for a in ADS],
  "precios": [dict(zip(["servicio","min","max","unidad","fuente"], p)) for p in PRECIOS],
  "legal": [dict(zip(["tema","regla","implicacion","fuente"], l)) for l in LEGAL],
  # Registros oficiales consultados en vivo desde el navegador (ArcGIS REST).
  "registros": REGISTROS,
  "coordenadas": {k: list(v) for k, v in COORDS.items()},
}
with open(out, "w", encoding="utf-8") as f:
    f.write("/* Generado por discover/research/scripts/export_funnel26.py desde el estudio de mercado.\n"
            "   No editar a mano: se regenera. Cada cifra lleva su fuente. */\n")
    f.write("export const DATOS = {\n")
    partes = []
    for k, v in datos.items():
        if isinstance(v, list):
            filas = ",\n".join("  " + json.dumps(x, ensure_ascii=False) for x in v)
            partes.append(f" {json.dumps(k)}: [\n{filas}\n ]")
        else:
            partes.append(f" {json.dumps(k)}: {json.dumps(v, ensure_ascii=False)}")
    f.write(",\n".join(partes) + "\n}\n")
print(out, len(pros), "prospectos")
